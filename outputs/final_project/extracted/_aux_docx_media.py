import zipfile
from docx import Document

SRC = r'inputs/final_project/proyecto/Anexos_CAIO vf.docx'
z = zipfile.ZipFile(SRC)
open('outputs/final_project/extracted/render/docx_image1.png', 'wb').write(z.read('word/media/image1.png'))
d = Document(SRC)
for k, p in enumerate(d.paragraphs, 1):
    if p._p.xpath('.//w:drawing'):
        print('drawing at paragraph', k, repr(p.text))
print('sections', len(d.sections))
for s in d.sections:
    print('header:', [p.text for p in s.header.paragraphs], 'footer:', [p.text for p in s.footer.paragraphs])
print('comments part:', [n for n in z.namelist() if 'comment' in n])
body = d.element.body
order = []
ti = pi = 0
for child in body.iterchildren():
    tag = child.tag.rsplit('}', 1)[-1]
    if tag == 'p':
        pi += 1
        order.append(f'p{pi}')
    elif tag == 'tbl':
        ti += 1
        order.append(f'T{ti}')
print(' '.join(order))
