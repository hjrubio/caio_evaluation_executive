#!/usr/bin/env python3
"""Inventario y extracción preliminar multimodal, sin afirmar comprensión visual ni semántica."""
import argparse, csv, hashlib, io, json, re, sys, zipfile
from pathlib import Path
from datetime import datetime, timezone
from xml.etree import ElementTree as ET

EXTS={'.doc','.docx','.gdoc','.pdf','.pptx','.ppt','.xls','.xlsx','.csv','.tsv','.txt','.md','.png','.jpg','.jpeg','.webp','.gif','.bmp','.tiff','.tif','.svg'}
IMAGES={'.png','.jpg','.jpeg','.webp','.gif','.bmp','.tiff','.tif','.svg'}
OFFICE_LEGACY={'.doc','.ppt','.xls','.gdoc'}
MAX_CHARS=1_000_000

def text_decode(raw):
    for enc in ('utf-8-sig','utf-16','cp1252','latin-1'):
        try:return raw.decode(enc),enc
        except UnicodeError:pass
    return raw.decode('utf-8',errors='replace'),'utf-8-replace'

def xml_text(xml, ending=' '):
    root=ET.fromstring(xml)
    return ending.join(t.text for t in root.iter() if t.tag.rsplit('}',1)[-1] in ('t','v') and t.text)

def extract(path):
    e=path.suffix.lower(); warnings=[]; units=[]; status='EXTRACTED_TEXT'; visual=False
    try:
        if e in ('.txt','.md'):
            txt,enc=text_decode(path.read_bytes()); units=[{'locator':'line:1','text':txt}]
            if enc!='utf-8-sig':warnings.append('codificación detectada: '+enc)
        elif e in ('.csv','.tsv'):
            txt,enc=text_decode(path.read_bytes()); dial=csv.excel_tab if e=='.tsv' else csv.excel
            rows=list(csv.reader(io.StringIO(txt),dialect=dial))
            for i,row in enumerate(rows,1): units.append({'locator':f'row:{i}','text':json.dumps(row,ensure_ascii=False)})
            warnings.append('Revisar encabezados, unidades, tipos de datos y fórmulas externas; los CSV no preservan formato.')
        elif e=='.docx':
            from docx import Document
            d=Document(path)
            for i,p in enumerate(d.paragraphs,1):
                if p.text.strip():units.append({'locator':f'paragraph:{i}','text':p.text})
            for t,table in enumerate(d.tables,1):
                for r,row in enumerate(table.rows,1):
                    units.append({'locator':f'table:{t}/row:{r}','text':' | '.join(c.text for c in row.cells)})
            visual=bool(d.inline_shapes)
            if visual:warnings.append('Contiene imágenes u objetos: validar visualmente, extracción textual insuficiente.')
            warnings.append('Revisar encabezados/pies, comentarios, cuadros de texto y diseños complejos en vista original.')
        elif e=='.pdf':
            import fitz
            with fitz.open(path) as pdf:
                for pno,page in enumerate(pdf,1):
                    txt=page.get_text('text')
                    units.append({'locator':f'page:{pno}','text':txt})
                    if page.get_images(full=True) or len(txt.strip())<30:visual=True
            if visual:warnings.append('Hay páginas gráficas, escaneadas o con poco texto: revisión visual requerida; OCR opcional con validación.')
            warnings.append('Validar tablas, figuras y disposición visual del PDF.')
        elif e=='.pptx':
            from pptx import Presentation
            prs=Presentation(path)
            for s,slide in enumerate(prs.slides,1):
                chunks=[]
                for shape in slide.shapes:
                    if shape.has_text_frame: chunks.append(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:chunks.append(' | '.join(c.text for c in row.cells))
                    if shape.shape_type==13 or shape.has_chart: visual=True
                units.append({'locator':f'slide:{s}','text':'\n'.join(chunks)})
            visual=True
            warnings.append('Revisión visual de todas las diapositivas obligatoria: gráficos, disposición, imágenes y notas pueden no quedar reflejados.')
        elif e=='.xlsx':
            visual=True
            ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            with zipfile.ZipFile(path) as z:
                shared=[]
                if 'xl/sharedStrings.xml' in z.namelist():
                    root=ET.fromstring(z.read('xl/sharedStrings.xml'))
                    shared=[''.join(x.text or '' for x in si.findall('.//m:t',ns)) for si in root.findall('m:si',ns)]
                sheets=sorted((n for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n)),key=lambda n:int(re.search(r'sheet(\d+)',n).group(1)))
                for s,name in enumerate(sheets,1):
                    root=ET.fromstring(z.read(name))
                    for row in root.findall('.//m:sheetData/m:row',ns):
                        cells=[]
                        for c in row.findall('m:c',ns):
                            v=c.find('m:v',ns); f=c.find('m:f',ns); inline=c.find('m:is',ns)
                            val=v.text if v is not None and v.text else ''
                            if c.get('t')=='s' and val.isdigit() and int(val)<len(shared):val=shared[int(val)]
                            if inline is not None:val=''.join(t.text or '' for t in inline.findall('.//m:t',ns))
                            if f is not None:val=f'FORMULA({f.text}) cached={val}'
                            if val:cells.append(f'{c.get("r")}={val}')
                        if cells:units.append({'locator':f'sheet:{s}/row:{row.get("r")}', 'text':' | '.join(cells)})
                if len(sheets)==0:warnings.append('No se encontraron hojas XML legibles.')
            warnings.append('Valores calculados no recalculados; validar fórmulas, tablas dinámicas, gráficos, hojas ocultas, nombres y formato en la planilla original.')
        elif e in IMAGES:
            from PIL import Image
            if e=='.svg':
                txt,_=text_decode(path.read_bytes()); units=[{'locator':'image:1','text':'SVG XML (no sustituye revisión visual): '+re.sub(r'<[^>]+>',' ',txt)[:3000]}]
            else:
                with Image.open(path) as im:
                    units=[{'locator':'image:1','text':f'Imagen {im.width}x{im.height}, modo {im.mode}; pendiente análisis visual por el agente multimodal.'}]
            visual=True; status='PENDING_VISUAL_REVIEW'
            warnings.append('La lectura visual por un modelo multimodal es obligatoria; no suponer contenido desde metadatos. No OCR como única evidencia.')
        elif e in OFFICE_LEGACY:
            status='UNSUPPORTED_NEEDS_CONVERSION'
            units=[{'locator':'file:1','text':'Convertir con herramienta confiable a DOCX/PPTX/XLSX o exportar Google Docs antes de evaluar.'}]
            warnings.append('Formato legacy o acceso externo no disponible; bloquea cobertura si es obligatorio.')
        else:status='UNSUPPORTED'
    except Exception as exc:
        status='EXTRACTION_FAILED';warnings.append(f'Error extrayendo: {type(exc).__name__}: {exc}')
    total=sum(len(u['text']) for u in units)
    if total>MAX_CHARS:
        status='PARTIAL_TRUNCATED'; warnings.append(f'Texto superior a {MAX_CHARS} caracteres: requiere extracción en lotes.')
        remaining=MAX_CHARS; reduced=[]
        for u in units:
            if remaining<=0:break
            reduced.append({'locator':u['locator'],'text':u['text'][:remaining]});remaining-=len(u['text'])
        units=reduced
    if status=='EXTRACTED_TEXT' and visual:status='PENDING_VISUAL_REVIEW'
    return status,visual,units,warnings

def main():
    p=argparse.ArgumentParser();p.add_argument('--docs',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    docs=a.docs.resolve();a.output.mkdir(parents=True,exist_ok=True)
    if not docs.is_dir():p.error('No existe --docs')
    entries=[]
    for i,f in enumerate(sorted(x for x in docs.rglob('*') if x.is_file() and x.suffix.lower() in EXTS),1):
        relative=f.relative_to(docs).as_posix();status,visual,units,warnings=extract(f)
        rec={'id':f'DOC-{i:03d}','path':relative,'format':f.suffix.lower(),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'size_bytes':f.stat().st_size,'status':status,'requires_visual_review':visual,'units_count':len(units),'warnings':warnings}
        entries.append(rec)
        (a.output/f'DOC-{i:03d}.json').write_text(json.dumps({'source':rec,'units':units},ensure_ascii=False,indent=2),encoding='utf-8')
    report={'created_utc':datetime.now(timezone.utc).isoformat(),'source_dir':str(docs),'sources':entries,'gate':'BLOCKED_DOCUMENT_REVIEW' if any(r['status']!='EXTRACTED_TEXT' for r in entries) else 'READY_FOR_SEMANTIC_REVIEW','disclaimer':'Extraer texto no equivale a entender ni verificar el contenido. El gate semántico y de contexto es independiente.'}
    (a.output/'multimodal_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'count':len(entries),'by_status':{s:sum(x['status']==s for x in entries) for s in sorted(set(x['status'] for x in entries))},'gate':report['gate']},ensure_ascii=False,indent=2))
    return 0 if entries else 2
if __name__=='__main__':sys.exit(main())
